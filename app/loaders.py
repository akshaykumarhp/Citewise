"""Document loaders.

Hand-written rather than imported from langchain-community (which is being sunset)
for two reasons:

1. Owning the loader means owning the metadata. Page numbers are what make real
   citations possible ("policy.pdf, p. 12") instead of vague ones ("policy.pdf").
2. Understanding what a Document actually is - text plus metadata - is the first
   thing worth learning about RAG.
"""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path

from langchain_core.documents import Document

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt", ".rst"}


def file_hash(path: Path) -> str:
    """Content hash, used by the ingestion manifest to skip unchanged files."""
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def _load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            # Scanned/image-only page. Phase 6 would OCR this.
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={"source": path.name, "path": str(path), "page": i},
            )
        )
    if not docs:
        logger.warning("%s produced no extractable text (scanned PDF?)", path.name)
    return docs


def _load_docx(path: Path) -> list[Document]:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Tables are frequently where the actual answers live in business documents.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": path.name, "path": str(path), "page": 1},
        )
    ]


def _load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": path.name, "path": str(path), "page": 1},
        )
    ]


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_text,
    ".markdown": _load_text,
    ".txt": _load_text,
    ".rst": _load_text,
}


def load_file(path: Path) -> list[Document]:
    """Load one file into a list of Documents (one per page for PDFs)."""
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise ValueError(f"Unsupported file type: {suffix}")
    docs = loader(path)
    digest = file_hash(path)
    for d in docs:
        d.metadata["file_hash"] = digest
    return docs


def discover_files(data_dir: Path) -> list[Path]:
    """Every supported file under data_dir, recursively, in stable order."""
    if not data_dir.exists():
        return []
    return sorted(
        p
        for p in data_dir.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        and not p.name.startswith(".")
    )
